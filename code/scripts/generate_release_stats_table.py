"""Write the release statistics table (manuscript table replacing fig. 2A).

Outputs, all in ``code/scripts/results/``:

``release_stats_all.csv``
    Every release tag, one row each, with model size and everything the
    phenotype scorer reports (confusion matrix, unscored categories,
    sensitivity, specificity). The full record; nothing is dropped.
``release_stats_table.csv`` / ``release_stats_table.docx``
    The manuscript table: draft (first release) vs current (latest release).
    Both written from the same dataframe, so they cannot disagree. ASM wants
    main-text tables as Word files.

Not run in CI: the numbers only change when a release is tagged, and each run
re-reads and re-scores every tagged model. Re-run it after a release, or after
editing ``data/known_growth_phenotypes.tsv`` (every release is re-scored
against the current table, so a new condition changes all the columns).

Needs the release tags locally (``git fetch --tags``)::

    python code/scripts/generate_release_stats_table.py
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from tools.release_stats import (  # noqa: E402
    STATISTIC_COLUMN,
    gather_release_stats,
    paper_table,
    section_rows,
)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(SCRIPT_DIR, "results")

# Same size as the known-growth-phenotypes table, so the two look alike in the
# submission. Production restyles main-text tables anyway.
DOCX_FONT_SIZE = Pt(9)

# Wide enough that "No growth observed, no growth predicted" stays on one line;
# Word's autofit otherwise splits the page evenly and wraps every outcome row.
LABEL_WIDTH = Inches(3.0)
VALUE_WIDTH = Inches(1.4)


def save_as_docx(table, path: str):
    """Write the manuscript table as Word: portrait, one header row.

    Section headings ("Model size", "Growth phenotypes") are bold and span the
    full width; statistic labels in those sections are indented under them.
    Numbers are right-aligned so the digits line up.

    No caption is written. The table legend is manuscript text and belongs to
    the author; add it in Word above the table.
    """
    document = Document()
    headings = set(section_rows(table))

    word_table = document.add_table(rows=1, cols=len(table.columns))
    word_table.style = "Table Grid"
    # Word only honours widths with autofit off and a width on every cell.
    word_table.autofit = False
    widths = [
        LABEL_WIDTH if name == STATISTIC_COLUMN else VALUE_WIDTH
        for name in table.columns
    ]
    for column, width in zip(word_table.columns, widths):
        column.width = width

    for cell, column_name, width in zip(
        word_table.rows[0].cells, table.columns, widths
    ):
        cell.width = width
        text = "" if column_name == STATISTIC_COLUMN else str(column_name)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = DOCX_FONT_SIZE
        if column_name != STATISTIC_COLUMN:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for position, record in enumerate(table.to_dict("records")):
        cells = word_table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
        if position in headings:
            merged = cells[0].merge(cells[-1])
            run = merged.paragraphs[0].add_run(record[STATISTIC_COLUMN])
            run.bold = True
            run.font.size = DOCX_FONT_SIZE
            continue
        for cell, column_name in zip(cells, table.columns):
            paragraph = cell.paragraphs[0]
            if column_name == STATISTIC_COLUMN:
                paragraph.paragraph_format.left_indent = Pt(10)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.add_run(str(record[column_name])).font.size = DOCX_FONT_SIZE

    document.save(path)


if __name__ == "__main__":
    os.makedirs(RESULTS_DIR, exist_ok=True)

    stats = gather_release_stats()
    stats.to_csv(os.path.join(RESULTS_DIR, "release_stats_all.csv"), index=False)

    table = paper_table(stats)
    table.to_csv(os.path.join(RESULTS_DIR, "release_stats_table.csv"), index=False)
    save_as_docx(table, os.path.join(RESULTS_DIR, "release_stats_table.docx"))

    print(table.to_string(index=False))
